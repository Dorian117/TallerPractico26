"""Lectores de documentos: UN lector por formato.

Cada lector recibe una ruta y devuelve un `DocumentoLeido`: una lista de
`Bloque`s (texto + pagina + seccion) mas metadatos del documento. El resto
del pipeline (chunking, embeddings...) no sabe de que formato vino el texto.

Por que separar "leer" de "trocear": cada formato sabe donde estan sus
encabezados (\\section en LaTeX, "Heading 1" en Word, #/## en Markdown, <h1>
en HTML, fuente grande en PDF). Si el lector entrega bloques ya agrupados por
seccion, el chunker puede cortar respetando la estructura del documento.
"""
from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass, field
from pathlib import Path


@dataclass
class Bloque:
    texto: str
    seccion: str = ""   # ruta de encabezados: "Ejemplos > Ejemplo 1: ..."
    pagina: int = 0     # 0 = el formato no tiene paginas (tex, md, html...)


@dataclass
class DocumentoLeido:
    archivo: str                       # nombre del archivo (la "fuente" que se cita)
    titulo: str                        # titulo legible del documento
    formato: str                       # tex | pdf | docx | md | txt | html
    sha256: str                        # para detectar archivos duplicados
    bloques: list[Bloque] = field(default_factory=list)
    advertencias: list[str] = field(default_factory=list)


def _sha256(ruta: Path) -> str:
    return hashlib.sha256(ruta.read_bytes()).hexdigest()


# ===========================================================================
# LaTeX (.tex, y .txt que contiene codigo LaTeX)
# ===========================================================================
# Entornos "protegidos": su contenido es matematica o estructura que no debe
# alterarse al limpiar. Las transformaciones de texto solo se aplican FUERA.
_ENTORNOS_MATE = (
    "align", "align*", "equation", "equation*", "gather", "gather*",
    "multline", "multline*", "eqnarray", "eqnarray*", "displaymath", "math",
)
_RE_PROTEGIDO = re.compile(
    r"\$\$.*?\$\$"                                   # $$ ... $$
    r"|\\\[.*?\\\]"                                  # \[ ... \]
    + "".join(
        rf"|\\begin\{{{re.escape(e)}\}}.*?\\end\{{{re.escape(e)}\}}"
        for e in _ENTORNOS_MATE
    ),
    re.S,
)
_RE_SECCION = re.compile(r"^\\(section|subsection|subsubsection)\*?\{", re.M)
_NIVEL = {"section": 1, "subsection": 2, "subsubsection": 3}


def _cierre_llave(texto: str, abre: int) -> int:
    """Indice de la '}' que cierra la '{' situada en `abre` (respeta anidamiento)."""
    prof = 0
    for i in range(abre, len(texto)):
        c = texto[i]
        if c == "\\":          # salta el caracter escapado (\{ \} \\)
            continue
        if c == "{" and (i == 0 or texto[i - 1] != "\\"):
            prof += 1
        elif c == "}" and texto[i - 1] != "\\":
            prof -= 1
            if prof == 0:
                return i
    return -1


def _argumento(texto: str, pos: int) -> tuple[str, int]:
    """Lee el argumento {...} que empieza en `pos`. Devuelve (contenido, fin)."""
    fin = _cierre_llave(texto, pos)
    if fin == -1:
        return texto[pos + 1:], len(texto)
    return texto[pos + 1:fin], fin + 1


def _reemplazar_comando(texto: str, nombre: str, plantilla: str) -> str:
    r"""Sustituye `\nombre{X}` por plantilla.format(X), con llaves anidadas."""
    patron = re.compile(r"\\" + re.escape(nombre) + r"\*?\s*\{")
    salida, i = [], 0
    while True:
        m = patron.search(texto, i)
        if not m:
            salida.append(texto[i:])
            return "".join(salida)
        contenido, fin = _argumento(texto, m.end() - 1)
        salida.append(texto[i:m.start()])
        salida.append(plantilla.format(contenido))
        i = fin


def _nodos_tikz(bloque: str) -> str:
    r"""Un diagrama TikZ no es texto, pero sus \node{...} si lo son (pasos de
    un algoritmo, listas de estrategias...). Se extrae el texto de cada nodo
    en orden en lugar de perderlo."""
    partes, i = [], 0
    for m in re.finditer(r"\\node\s*(?:\[[^\]]*\])?\s*(?:\([^)]*\))?\s*(?:at\s*\([^)]*\))?\s*\{", bloque):
        if m.start() < i:
            continue
        contenido, fin = _argumento(bloque, m.end() - 1)
        i = fin
        partes.append(contenido.strip())
    return "Diagrama:\n" + "\n".join(f"- {p}" if "\n" not in p else p for p in partes)


def _tabla_a_texto(entorno: str) -> str:
    """tabular -> filas 'a | b | c' (las formulas de las celdas se conservan)."""
    cuerpo = re.sub(r"^\\begin\{tabular\}\s*\{[^}]*\}", "", entorno.strip())
    cuerpo = re.sub(r"\\end\{tabular\}\s*$", "", cuerpo)
    cuerpo = re.sub(r"\\(hline|toprule|midrule|bottomrule)", "", cuerpo)
    filas = []
    for fila in re.split(r"\\\\(?:\[[^\]]*\])?", cuerpo):
        celdas = [c.strip() for c in fila.split("&")]
        if any(celdas):
            filas.append("| " + " | ".join(celdas) + " |")
    return "\n".join(filas)


def _limpiar_fuera_de_mate(texto: str) -> str:
    """Transformaciones de TEXTO (no tocan formulas)."""
    # tikz: se conserva el texto de los nodos
    texto = re.sub(
        r"\\begin\{tikzpicture\}.*?\\end\{tikzpicture\}",
        lambda m: _nodos_tikz(m.group(0)), texto, flags=re.S,
    )
    # tablas
    texto = re.sub(
        r"\\begin\{tabular\}.*?\\end\{tabular\}",
        lambda m: _tabla_a_texto(m.group(0)), texto, flags=re.S,
    )
    # entornos de maquetacion: se quitan las marcas, se conserva el contenido
    texto = re.sub(
        r"\\(?:begin|end)\{(?:center|multicols|tcolorbox|minipage|figure|table|document)\}"
        r"(?:\{[^}]*\})?(?:\[[^\]]*\])?",
        "", texto,
    )
    # listas -> markdown ("- "); se quita la marca del entorno y su argumento
    texto = re.sub(r"\\(?:begin|end)\{(?:itemize|enumerate)\}(?:\[[^\]]*\])?", "", texto)
    texto = re.sub(r"\\item(?:\[[^\]]*\])?\s*", "- ", texto)
    # comandos puramente tipograficos
    texto = re.sub(
        r"\\(?:newpage|clearpage|noindent|medskip|bigskip|smallskip|hfill|centering|"
        r"maketitle|raggedright|small|large|Large|LARGE|huge|Huge|normalsize)\b\s*",
        "", texto,
    )
    texto = re.sub(r"\\(?:vspace|hspace|setlength|renewcommand|rule)\*?(?:\{[^}]*\})+(?:\[[^\]]*\])?", "", texto)
    # notas al pie y enfasis
    texto = _reemplazar_comando(texto, "footnote", " (nota: {})")
    texto = _reemplazar_comando(texto, "textbf", "**{}**")
    texto = _reemplazar_comando(texto, "textit", "*{}*")
    texto = _reemplazar_comando(texto, "emph", "*{}*")
    texto = _reemplazar_comando(texto, "underline", "{}")
    # saltos de linea de texto, simbolos escapados, espacio duro
    texto = re.sub(r"\\\\(?:\[[^\]]*\])?", "\n", texto)
    texto = texto.replace("\\%", "%").replace("\\&", "&").replace("~", " ")
    # sangria heredada del codigo fuente (solo tras un salto de linea, para no
    # pegar palabras cuando el tramo empieza justo despues de una formula)
    texto = re.sub(r"(?<=\n)[ \t]+", "", texto)
    return texto


def limpiar_tex(fuente: str) -> str:
    """LaTeX del cuerpo -> texto limpio que conserva la matematica en LaTeX.

    Se hace asi porque el LLM (y el estudiante) necesitan ver las formulas en
    LaTeX, pero no el ruido de maquetacion (\\vspace, \\hline, comentarios).
    """
    # comentarios (un % sin escapar hasta fin de linea)
    fuente = re.sub(r"(?<!\\)%.*", "", fuente)
    # lineas '---' sueltas (separadores que no son LaTeX valido)
    fuente = re.sub(r"^\s*-{3,}\s*$", "", fuente, flags=re.M)

    # se separan las formulas protegidas del resto y solo se limpia el resto
    salida, i = [], 0
    for m in _RE_PROTEGIDO.finditer(fuente):
        salida.append(_limpiar_fuera_de_mate(fuente[i:m.start()]))
        salida.append(m.group(0))
        i = m.end()
    salida.append(_limpiar_fuera_de_mate(fuente[i:]))
    texto = "".join(salida)

    texto = re.sub(r"[ \t]+\n", "\n", texto)
    texto = re.sub(r"\n{3,}", "\n\n", texto)
    return texto.strip()


def _titulo_tex(fuente: str, respaldo: str) -> str:
    # 1) \title{\textbf{Calculo II: X}\\ ...}
    m = re.search(r"\\title\s*\{", fuente)
    if m:
        contenido, _ = _argumento(fuente, m.end() - 1)
        t = re.search(r"\\textbf\s*\{", contenido)
        if t:
            return _argumento(contenido, t.end() - 1)[0].strip()
    # 2) encabezado propio: {\LARGE \textbf{TITULO}}
    m = re.search(r"\\LARGE\s*\\textbf\s*\{", fuente)
    if m:
        return _argumento(fuente, m.end() - 1)[0].strip()
    return respaldo


def es_latex(texto: str) -> bool:
    """Detecta LaTeX por contenido: la extension .txt no es de fiar."""
    return "\\documentclass" in texto[:2000] or "\\begin{document}" in texto


def leer_tex_texto(fuente: str, archivo: str, sha: str, formato: str = "tex") -> DocumentoLeido:
    titulo = _titulo_tex(fuente, Path(archivo).stem)
    ini = fuente.find("\\begin{document}")
    fin = fuente.find("\\end{document}")
    cuerpo = fuente[ini + len("\\begin{document}"): fin if fin != -1 else None]

    # El primer bloque centrado con \rule o la institucion es el membrete
    # (universidad, facultad, docente): no es contenido del curso.
    m = re.match(r"\s*\\begin\{center\}.*?\\end\{center\}", cuerpo, re.S)
    if m and ("\\rule" in m.group(0) or "FUNDACI" in m.group(0).upper()):
        cuerpo = cuerpo[m.end():]

    doc = DocumentoLeido(archivo=archivo, titulo=titulo, formato=formato, sha256=sha)

    # Se parte por \section / \subsection / \subsubsection manteniendo una
    # pila de encabezados: la "seccion" de un bloque es la ruta completa.
    marcas = []
    for m in _RE_SECCION.finditer(cuerpo):
        nombre, fin_arg = _argumento(cuerpo, m.end() - 1)
        marcas.append((m.start(), fin_arg, _NIVEL[m.group(1)], nombre.strip()))
    marcas.append((len(cuerpo), len(cuerpo), 0, ""))

    pila: list[tuple[int, str]] = []
    primer = cuerpo[: marcas[0][0]]
    tramos = [("Introduccion", primer)] if primer.strip() else []
    for (ini_m, fin_arg, nivel, nombre), (sig, _, _, _) in zip(marcas, marcas[1:]):
        while pila and pila[-1][0] >= nivel:
            pila.pop()
        pila.append((nivel, nombre))
        tramos.append((" > ".join(n for _, n in pila), cuerpo[fin_arg:sig]))

    for seccion, crudo in tramos:
        texto = limpiar_tex(crudo)
        if texto:  # una seccion sin cuerpo propio (solo subsecciones) no genera bloque
            doc.bloques.append(Bloque(texto=texto, seccion=seccion))
    return doc


def leer_tex(ruta: Path) -> DocumentoLeido:
    return leer_tex_texto(ruta.read_text(encoding="utf-8"), ruta.name, _sha256(ruta))


# ===========================================================================
# TXT / Markdown
# ===========================================================================
def leer_txt(ruta: Path) -> DocumentoLeido:
    """.txt: si en realidad es LaTeX se delega en el lector LaTeX (el corpus
    del curso son archivos .txt con codigo LaTeX); si no, es texto plano."""
    texto = ruta.read_text(encoding="utf-8")
    if es_latex(texto):
        return leer_tex_texto(texto, ruta.name, _sha256(ruta), formato="tex")
    return DocumentoLeido(
        archivo=ruta.name, titulo=ruta.stem, formato="txt", sha256=_sha256(ruta),
        bloques=[Bloque(texto=texto.strip())] if texto.strip() else [],
    )


def leer_md(ruta: Path) -> DocumentoLeido:
    texto = ruta.read_text(encoding="utf-8")
    doc = DocumentoLeido(archivo=ruta.name, titulo=ruta.stem, formato="md", sha256=_sha256(ruta))
    pila: list[tuple[int, str]] = []
    buf: list[str] = []

    def volcar():
        cuerpo = "\n".join(buf).strip()
        if cuerpo:
            doc.bloques.append(Bloque(texto=cuerpo, seccion=" > ".join(n for _, n in pila)))
        buf.clear()

    en_codigo = False
    for linea in texto.splitlines():
        if linea.strip().startswith("```"):
            en_codigo = not en_codigo  # un '#' dentro de un bloque de codigo no es encabezado
        m = None if en_codigo else re.match(r"^(#{1,6})\s+(.*\S)\s*$", linea)
        if m:
            volcar()
            nivel = len(m.group(1))
            while pila and pila[-1][0] >= nivel:
                pila.pop()
            pila.append((nivel, m.group(2)))
            if nivel == 1 and doc.titulo == ruta.stem:
                doc.titulo = m.group(2)
        else:
            buf.append(linea)
    volcar()
    return doc


# ===========================================================================
# HTML
# ===========================================================================
def leer_html(ruta: Path) -> DocumentoLeido:
    from bs4 import BeautifulSoup, NavigableString, Tag

    soup = BeautifulSoup(ruta.read_text(encoding="utf-8"), "html.parser")
    for t in soup(["script", "style", "nav", "footer", "header", "noscript"]):
        t.decompose()
    titulo = soup.title.get_text(" ", strip=True) if soup.title else ruta.stem
    doc = DocumentoLeido(archivo=ruta.name, titulo=titulo, formato="html", sha256=_sha256(ruta))

    encabezados = {f"h{i}": i for i in range(1, 7)}
    saltar: set[int] = set()
    pila: list[tuple[int, str]] = []
    buf: list[str] = []

    def volcar():
        cuerpo = re.sub(r"\n{3,}", "\n\n", "".join(buf)).strip()
        if cuerpo:
            doc.bloques.append(Bloque(texto=cuerpo, seccion=" > ".join(n for _, n in pila)))
        buf.clear()

    for el in (soup.body or soup).descendants:
        if isinstance(el, Tag):
            if el.name in encabezados:
                volcar()
                nivel = encabezados[el.name]
                while pila and pila[-1][0] >= nivel:
                    pila.pop()
                pila.append((nivel, el.get_text(" ", strip=True)))
                saltar.update(id(s) for s in el.descendants)
            elif el.name in ("p", "li", "br", "div", "tr", "pre", "blockquote"):
                buf.append("\n")
        elif isinstance(el, NavigableString) and id(el) not in saltar:
            buf.append(str(el))
    volcar()
    return doc


# ===========================================================================
# DOCX
# ===========================================================================
def leer_docx(ruta: Path) -> DocumentoLeido:
    import docx

    d = docx.Document(str(ruta))
    doc = DocumentoLeido(archivo=ruta.name, titulo=ruta.stem, formato="docx", sha256=_sha256(ruta))
    pila: list[tuple[int, str]] = []
    buf: list[str] = []

    def volcar():
        cuerpo = "\n".join(buf).strip()
        if cuerpo:
            doc.bloques.append(Bloque(texto=cuerpo, seccion=" > ".join(n for _, n in pila)))
        buf.clear()

    n_ecuaciones = 0
    for p in d.paragraphs:
        # python-docx ignora las ecuaciones de Word (OMML): se rescata su texto
        # lineal (m:t) para no perderlas en silencio, aunque sin estructura.
        mate = [t.text for t in p._p.iter() if t.tag.endswith("}t") and "math" in t.tag]
        n_ecuaciones += bool(mate)
        texto = (p.text + (" " + "".join(mate) if mate else "")).strip()
        if not texto:
            continue
        estilo = p.style.name if p.style is not None else ""
        m = re.match(r"Heading (\d)", estilo)
        if m:
            volcar()
            nivel = int(m.group(1))
            while pila and pila[-1][0] >= nivel:
                pila.pop()
            pila.append((nivel, texto))
        else:
            buf.append(texto)
    for tabla in d.tables:
        for fila in tabla.rows:
            buf.append("| " + " | ".join(c.text.strip() for c in fila.cells) + " |")
    volcar()
    if n_ecuaciones:
        doc.advertencias.append(
            f"{n_ecuaciones} parrafos con ecuaciones OMML: solo se rescato su texto lineal, "
            "no LaTeX. Preferible exportar a .tex."
        )
    return doc


# ===========================================================================
# PDF
# ===========================================================================
def _calidad_pdf(paginas: list[str]) -> list[str]:
    """Senales objetivas de extraccion deformada (formulas rotas)."""
    texto = "\n".join(paginas)
    avisos = []
    lineas = [l for l in texto.splitlines() if l.strip()]
    cortas = sum(1 for l in lineas if len(l.strip()) <= 2) / max(len(lineas), 1)
    if cortas > 0.15:
        avisos.append(f"{100*cortas:.0f}% de lineas de <=2 caracteres: formulas probablemente fragmentadas")
    if re.search(r"[\U0001D400-\U0001D7FF]\s?[!\-]", texto):
        avisos.append("patrones tipo 'x!' / 'x-' tras letras matematicas: exponentes mal extraidos")
    if re.search(r"[\ue000-\uf8ff\ufffd]", texto) or "(cid:" in texto:
        avisos.append("caracteres de area privada / reemplazo: fuente sin ToUnicode")
    vacias = sum(1 for p in paginas if len(p.strip()) < 40)
    if vacias:
        avisos.append(f"{vacias} paginas casi sin texto (escaneadas o solo imagenes)")
    return avisos


def leer_pdf(ruta: Path) -> DocumentoLeido:
    """PDF con pymupdf. La 'seccion' se infiere por tamano de fuente: una
    linea corta con letra claramente mayor que la mediana es un encabezado."""
    import pymupdf

    doc = DocumentoLeido(archivo=ruta.name, titulo=ruta.stem, formato="pdf", sha256=_sha256(ruta))
    paginas_texto, seccion = [], ""
    with pymupdf.open(ruta) as pdf:
        if pdf.metadata and pdf.metadata.get("title"):
            doc.titulo = pdf.metadata["title"].strip() or doc.titulo
        for n, pagina in enumerate(pdf, start=1):
            lineas = []  # (texto, tamano)
            for b in pagina.get_text("dict")["blocks"]:
                for l in b.get("lines", []):
                    t = "".join(s["text"] for s in l["spans"]).strip()
                    if t:
                        lineas.append((t, max(s["size"] for s in l["spans"])))
            tamanos = sorted(s for _, s in lineas)
            mediana = tamanos[len(tamanos) // 2] if tamanos else 0
            cuerpo = []
            for t, s in lineas:
                if mediana and s >= mediana * 1.2 and len(t) <= 100:
                    seccion = t
                else:
                    cuerpo.append(t)
            texto = "\n".join(cuerpo)
            paginas_texto.append(texto)
            if texto.strip():
                doc.bloques.append(Bloque(texto=texto, seccion=seccion, pagina=n))
    doc.advertencias.extend(_calidad_pdf(paginas_texto))
    return doc


# ===========================================================================
# Despacho por formato
# ===========================================================================
LECTORES = {
    ".tex": leer_tex,
    ".txt": leer_txt,
    ".md": leer_md,
    ".markdown": leer_md,
    ".html": leer_html,
    ".htm": leer_html,
    ".docx": leer_docx,
    ".pdf": leer_pdf,
}


def leer_documento(ruta: Path) -> DocumentoLeido:
    lector = LECTORES.get(ruta.suffix.lower())
    if lector is None:
        raise ValueError(f"Formato no soportado: {ruta.name}")
    return lector(ruta)
